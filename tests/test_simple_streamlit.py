from __future__ import annotations

from io import BytesIO
from typing import ClassVar

import pymupdf
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from grounded_docparse import pipeline, universal
from grounded_docparse.agentic import DocumentAgent
from grounded_docparse.batch import build_batch_documents
from grounded_docparse.config import LUNA_MODEL
from grounded_docparse.models import (
    AgenticAnalysis,
    AgentUsage,
    Block,
    Document,
    Page,
    ParseMetadata,
    ParseResult,
    ProgressEvent,
    RunUsage,
    SchemaField,
    StoredSchema,
    VerificationState,
)
from grounded_docparse.native import (
    NativeDocument,
    NativeParseResult,
    PageRoute,
    ProcessingType,
    SourceFormat,
    SourceUnit,
)
from grounded_docparse.render import build_elements, render_agentic_document
from grounded_docparse.schema_store import SchemaStore
from grounded_docparse.universal import PdfInspection
from grounded_docparse.workspace_store import WorkspaceStore


@pytest.fixture(autouse=True)
def isolated_studio_database(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv("DOCPARSE_STUDIO_DB_PATH", str(tmp_path / "studio.sqlite3"))
    monkeypatch.setattr(
        universal,
        "inspect_pdf_content",
        lambda _data: PdfInspection(
            pdf_type="scanned",
            page_count=1,
            pages_needing_ocr=frozenset({1}),
        ),
    )


def _select_scanned(app: AppTest, filename: str) -> AppTest:
    return next(
        item
        for item in app.selectbox
        if item.label == f"Processing type · {filename}"
    ).select("Scanned PDF").run(timeout=20)


def test_studio_allows_glm_without_openai_environment(
    monkeypatch, simple_pdf: bytes
) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("OPENAI_BASE_URL", raising=False)

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)

    assert not app.exception
    assert len(app.file_uploader) == 1
    assert app.sidebar
    app = app.file_uploader[0].upload(
        "notice.pdf", simple_pdf, "application/pdf"
    ).run(timeout=20)
    parse = next(button for button in app.button if button.label == "Parse document")
    assert parse.disabled is True
    app = _select_scanned(app, "notice.pdf")
    parse = next(button for button in app.button if button.label == "Parse document")
    assert parse.disabled is False
    assert next(item for item in app.selectbox if item.label == "ADE mode").value == "Fast"
    model = next(
        item for item in app.selectbox if item.label == "Document extraction model"
    )
    assert model.options == ["GLM-OCR", "PaddleOCR-VL-1.6"]
    assert model.value == "GLM-OCR"
    assert not next(
        toggle.value
        for toggle in app.toggle
        if toggle.label == "Enhance with gpt-5.6-luna"
    )
    recovery = next(
        toggle
        for toggle in app.toggle
        if toggle.label == "Enable visual recovery on hard regions"
    )
    assert recovery.disabled is True
    assert recovery.help == (
        "Uses medium-effort Luna vision on prioritized hard regions. The budget "
        "scales from 8 to 64 crops by document length, with at most 3 per page; "
        "local GLM recovery runs first."
    )
    assert [tab.label for tab in app.tabs] == [
        "Overview",
        "Markdown",
        "Annotated PDF",
        "Layout Tree",
    ]


def test_stale_session_result_is_discarded(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    app = AppTest.from_file("streamlit_app.py")
    app.session_state["result"] = object()
    app.session_state["result_source_hash"] = "old"
    app.session_state["result_version"] = "old-contract"
    app.session_state["session_usage"] = RunUsage(
        calls=[
            AgentUsage(
                agent="old",
                model=LUNA_MODEL,
                input_tokens=100,
                output_tokens=10,
            )
        ]
    )

    app.run(timeout=20)

    assert app.session_state["result"] is None
    assert app.session_state["result_source_hash"] is None
    assert app.session_state["result_version"] == "4.6.0"
    assert app.session_state["session_usage"].calls == []


def test_legacy_session_usage_defaults_missing_cached_tokens_to_zero(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    app = AppTest.from_file("streamlit_app.py").run(timeout=20)

    class LegacyUsage:
        calls: ClassVar = [
            type(
                "LegacyAgentUsage",
                (),
                {
                    "model": LUNA_MODEL,
                    "input_tokens": 100,
                    "output_tokens": 10,
                },
            )()
        ]

    app.session_state["session_usage"] = LegacyUsage()
    app = app.run(timeout=20)

    assert not app.exception
    metrics = {item.label: item.value for item in app.metric}
    assert metrics["Cached input"] == "0"
    assert metrics["Estimated cost"] == "$0.000032"


def test_studio_shows_default_luna_destination(monkeypatch) -> None:
    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.delenv("OPENAI_BASE_URL", raising=False)

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)

    assert any(
        "Luna destination: OpenAI default endpoint" in item.value
        for item in app.get("caption")
    )


def test_ocr_model_selection_updates_the_active_ui_engine(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    app = AppTest.from_file("streamlit_app.py").run(timeout=20)

    model = next(
        item for item in app.selectbox if item.label == "Document extraction model"
    )
    app = model.select("PaddleOCR-VL-1.6").run(timeout=20)

    assert not app.exception
    assert any(
        "Powered by PaddleOCR-VL-1.6 + gpt-5.6-luna" in item.value
        for item in app.get("caption")
    )


def test_studio_hides_custom_luna_destination(monkeypatch) -> None:
    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setenv("OPENAI_BASE_URL", "https://us.api.openai.com/v1")

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)

    assert not any("custom endpoint" in item.value for item in app.warning)
    assert not any("Luna destination" in item.value for item in app.get("caption"))


def test_ade_presets_default_fast_and_allow_full_or_custom(monkeypatch) -> None:
    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    app = AppTest.from_file("streamlit_app.py").run(timeout=20)

    mode = next(item for item in app.selectbox if item.label == "ADE mode")
    assert mode.value == "Fast"
    app = mode.select("Full").run(timeout=20)
    toggles = {item.label: item for item in app.toggle}
    assert toggles["Enhance with gpt-5.6-luna"].value is True
    assert toggles["Classify document type"].value is True
    assert toggles["Generate table of contents"].value is True
    assert toggles["Enable document chat"].value is False
    assert toggles["Enable visual recovery on hard regions"].value is True

    app = toggles["Generate table of contents"].set_value(False).run(timeout=20)
    assert next(item for item in app.selectbox if item.label == "ADE mode").value == "Custom"


def test_studio_shows_results_and_only_requested_tools(
    monkeypatch, simple_pdf: bytes, tmp_path
) -> None:
    schema_database = tmp_path / "studio.sqlite3"
    monkeypatch.setenv("DOCPARSE_STUDIO_DB_PATH", str(schema_database))
    document = Document(
        source_name="notice.pdf",
        source_sha256="a" * 64,
        pages=[
            Page(
                number=1,
                width=612,
                height=792,
                blocks=[
                    Block(
                        id="p1-b1",
                        type="heading",
                        text="Public notice",
                        reading_order=0,
                        verification=VerificationState.VERIFIED,
                    )
                ],
            )
        ],
    )

    class FakeParser:
        def __init__(self, config=None):
            assert config.ocr_engine.value == "glm-ocr"

        def parse(
            self,
            _data,
            _name,
            progress_callback=None,
            *,
            refine_markdown=True,
            visual_recovery=True,
        ):
            assert visual_recovery is True
            del refine_markdown
            if progress_callback is not None:
                for stage in ("layout", "recognize", "assemble", "annotate", "complete"):
                    progress_callback(
                        ProgressEvent(
                            stage=stage,
                            current=1,
                            total=1,
                            message=stage,
                        )
                    )
            rendered = render_agentic_document(document)
            return ParseResult(
                document=document,
                markdown=rendered.markdown,
                json=rendered.json,
                input_tokens=1_234,
                output_tokens=56,
                annotated_pdf=simple_pdf,
                usage=RunUsage(
                    calls=[
                        AgentUsage(
                            agent="visual_recovery",
                            model=LUNA_MODEL,
                            input_tokens=1_234,
                            cached_input_tokens=234,
                            output_tokens=56,
                        )
                    ]
                ),
                trace=[],
                elements=build_elements(document, {"p1-b1"}),
                metadata=ParseMetadata(
                    visual_recovery_candidates=1,
                    visual_recovery_crops=1,
                    visual_recovery_region_ids=["p1-b1"],
                ),
            )

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "DocumentParser", FakeParser)
    monkeypatch.setattr(DocumentAgent, "analyze", lambda self, *args, **kwargs: AgenticAnalysis())

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app.file_uploader[0].upload(
        "notice.pdf", simple_pdf, "application/pdf"
    ).run(timeout=20)
    app = _select_scanned(app, "notice.pdf")
    app = next(
        button for button in app.button if button.label == "Parse document"
    ).click().run(timeout=20)

    assert not app.exception
    assert [tab.label for tab in app.tabs] == [
        "Overview",
        "Markdown",
        "Annotated PDF",
        "Extract",
        "Layout Tree",
    ]
    assert next(item for item in app.metric if item.label == "Recovered").value == "1"
    metrics = {item.label: item.value for item in app.metric}
    assert metrics == {
        "Input tokens": "1,234",
        "Cached input": "234",
        "Output tokens": "56",
        "Estimated cost": "$0.000272",
        "Pages": "1",
        "Regions": "1",
        "Tables": "0",
        "Figures": "0",
        "Time": "0.0s",
        "Recovered": "1",
    }
    assert not app.text_area
    assert [button.label for button in app.download_button] == [
        "Download Markdown",
        "Download annotated PDF",
        "Download Full JSON",
        "Download all outputs",
    ]
    assert any("$0.02/M cached input" in item.value for item in app.get("caption"))
    assert any(
        "Review confidence and cited source regions" in item.value
        for item in app.warning
    )

    app.session_state["studio_tab"] = "Markdown"
    app = app.run(timeout=20)
    assert app.session_state["session_usage"].input_tokens == 1_234
    preview = next(item for item in app.markdown if "Public notice" in item.value)
    assert preview.proto.allow_html is True

    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    assert not app.exception
    assert any("Define the field keys" in item.value for item in app.get("caption"))
    assert any(button.label == "Run extraction" for button in app.button) is False

    schema_uploader = next(
        item
        for item in app.file_uploader
        if item.label == "Import schema Markdown, CSV, or XLSX"
    )
    schema_uploader.upload(
        "invoice.md",
        b"# Invoice\n- invoice_number: Official invoice ID",
        "text/markdown",
    )
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    assert not app.exception
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    assert not app.exception
    assert app.session_state["schema_draft_name"] == "Invoice"
    assert app.session_state["schema_draft_fields"] == [
        {
            "name": "invoice_number",
            "description": "Official invoice ID",
            "type": "string",
        }
    ]
    assert any(button.label == "Run extraction" for button in app.button)

    schema_uploader = next(
        item
        for item in app.file_uploader
        if item.label == "Import schema Markdown, CSV, or XLSX"
    )
    schema_uploader.upload(
        "payments.csv",
        (
            b"Field name,Description,Type\n"
            b"total_amount,Final amount payable,number\n"
        ),
        "text/csv",
    )
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    assert app.session_state["schema_draft_name"] == "payments"
    assert app.session_state["schema_draft_fields"][0]["name"] == "total_amount"

    workbook = Workbook()
    workbook.active.append(["Field name", "Description", "Type"])
    workbook.active.append(["due_date", "Payment due date", "date"])
    output = BytesIO()
    workbook.save(output)
    schema_uploader = next(
        item
        for item in app.file_uploader
        if item.label == "Import schema Markdown, CSV, or XLSX"
    )
    schema_uploader.upload(
        "dates.xlsx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    assert app.session_state["schema_draft_name"] == "dates"
    assert app.session_state["schema_draft_fields"][0]["name"] == "due_date"
    assert SchemaStore(schema_database).list() == []

    SchemaStore(schema_database).save(
        StoredSchema(
            name="Invoice",
            fields=[
                SchemaField(
                    name="invoice_number",
                    description="Official invoice ID",
                )
            ],
        )
    )
    app.session_state["use_custom_routing"] = True
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    routing_uploader = next(
        item
        for item in app.file_uploader
        if item.label == "Import routing profile Markdown"
    )
    routing_uploader.upload(
        "medical-routing.md",
        (
            b"# Medical routing\n"
            b"- newauth [extract=Invoice]: Initial authorization request"
        ),
        "text/markdown",
    )
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    app.session_state["studio_tab"] = "Extract"
    app = app.run(timeout=20)
    assert not app.exception
    assert app.session_state["classifier_draft_name"] == "Medical routing"
    assert app.session_state["classifier_draft_categories"][0]["key"] == "newauth"
    assert any(button.label == "Classify forms" for button in app.button)

    app.session_state["studio_tab"] = "Layout Tree"
    app = app.run(timeout=20)
    assert any(button.label.startswith("1 · Heading") for button in app.button)
    assert any("Luna" in item.value for item in app.markdown)

    app.session_state.enable_chat = True
    app.session_state.chat_history = [
        {"role": "assistant", "content": "Grounded answer", "confidence": "high"}
    ]
    app = app.run(timeout=20)
    app.session_state["studio_tab"] = "Chat"
    app = app.run(timeout=20)
    assert any("Confidence: high" in item.value for item in app.get("caption"))


def test_page_range_parses_a_renumbered_pdf_subset(monkeypatch) -> None:
    source = pymupdf.open()
    source.new_page().insert_text((72, 72), "First page")
    source.new_page().insert_text((72, 72), "Second page")
    source_bytes = source.tobytes()
    source.close()
    captured_page_counts: list[int] = []
    captured_recovery: list[bool] = []

    class FakeParser:
        def __init__(self, config=None):
            assert config.ocr_engine.value == "glm-ocr"

        def parse(
            self,
            data,
            name,
            progress_callback=None,
            *,
            refine_markdown=True,
            visual_recovery=True,
        ):
            del refine_markdown
            captured_recovery.append(visual_recovery)
            with pymupdf.open(stream=data, filetype="pdf") as selected:
                captured_page_counts.append(selected.page_count)
            document = Document(
                source_name=name,
                source_sha256="b" * 64,
                pages=[Page(number=1, width=595, height=842)],
            )
            rendered = render_agentic_document(document)
            return ParseResult(
                document=document,
                markdown=rendered.markdown,
                json=rendered.json,
                input_tokens=0,
                output_tokens=0,
                annotated_pdf=data,
            )

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "DocumentParser", FakeParser)
    monkeypatch.setattr(DocumentAgent, "analyze", lambda self, *args, **kwargs: AgenticAnalysis())

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app.file_uploader[0].upload(
        "two-pages.pdf", source_bytes, "application/pdf"
    ).run(timeout=20)
    app = _select_scanned(app, "two-pages.pdf")
    app = next(box for box in app.checkbox if box.label == "Page range").check().run(
        timeout=20
    )
    app = next(item for item in app.number_input if item.label == "Start").set_value(
        2
    ).run(timeout=20)
    app = next(
        item
        for item in app.toggle
        if item.label == "Enable visual recovery on hard regions"
    ).set_value(False).run(timeout=20)
    app = next(
        button for button in app.button if button.label == "Parse document"
    ).click().run(timeout=20)

    assert not app.exception
    assert captured_page_counts == [1]
    assert captured_recovery == [False]
    assert app.session_state.overview_page == 1


def test_multiple_uploads_process_sequentially_and_only_process_new_files(
    monkeypatch, simple_pdf: bytes
) -> None:
    parsed_names: list[str] = []

    class FakeParser:
        def __init__(self, config=None):
            assert config.ocr_engine.value == "glm-ocr"

        def parse(self, data, name, **_kwargs):
            parsed_names.append(name)
            document = Document(
                source_name=name,
                source_sha256="c" * 64,
                pages=[Page(number=1, width=612, height=792)],
            )
            rendered = render_agentic_document(document)
            return ParseResult(
                document=document,
                markdown=rendered.markdown,
                json=rendered.json,
                input_tokens=10,
                output_tokens=1,
                annotated_pdf=data,
                usage=RunUsage(
                    calls=[
                        AgentUsage(
                            agent="visual_recovery",
                            model=LUNA_MODEL,
                            input_tokens=10,
                            output_tokens=1,
                        )
                    ]
                ),
            )

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "DocumentParser", FakeParser)
    monkeypatch.setattr(
        DocumentAgent, "analyze", lambda self, *args, **kwargs: AgenticAnalysis()
    )

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    uploader = app.file_uploader[0]
    uploader.upload("first.pdf", simple_pdf, "application/pdf")
    uploader.upload("second.pdf", simple_pdf + b"\n", "application/pdf")
    app = uploader.run(timeout=20)
    app = _select_scanned(app, "first.pdf")
    app = _select_scanned(app, "second.pdf")

    page_range = next(box for box in app.checkbox if box.label == "Page range")
    assert page_range.disabled is True
    app = next(
        button for button in app.button if button.label == "Process 2 documents"
    ).click().run(timeout=20)

    assert not app.exception
    assert parsed_names == ["first.pdf", "second.pdf"]
    assert [
        workspace["status"]
        for workspace in app.session_state["batch_workspaces"].values()
    ] == ["complete", "complete"]
    assert app.session_state["session_usage"].input_tokens == 20
    selector = next(
        item for item in app.selectbox if item.label == "Document results"
    )
    app = selector.select("second.pdf").run(timeout=20)
    assert app.session_state["result"].document.source_name == "second.pdf"

    uploader = app.file_uploader[0]
    uploader.set_value(
        [
            ("first.pdf", simple_pdf, "application/pdf"),
            ("second.pdf", simple_pdf + b"\n", "application/pdf"),
            ("third.pdf", simple_pdf + b"\n\n", "application/pdf"),
        ]
    )
    app = uploader.run(timeout=20)
    app = _select_scanned(app, "third.pdf")
    assert not app.exception
    button_labels = [button.label for button in app.button]
    assert "Process 3 documents" in button_labels, button_labels
    app = next(
        button for button in app.button if button.label == "Process 3 documents"
    ).click().run(timeout=20)

    assert not app.exception
    assert parsed_names == ["first.pdf", "second.pdf", "third.pdf"]
    assert app.session_state["session_usage"].input_tokens == 30
    assert any(
        button.label == "Download all outputs" for button in app.download_button
    )


def test_batch_continues_after_failure_and_retry_skips_completed_document(
    monkeypatch, simple_pdf: bytes
) -> None:
    parsed_names: list[str] = []
    bad_attempts = 0

    class FakeParser:
        def __init__(self, config=None):
            assert config.ocr_engine.value == "glm-ocr"

        def parse(self, data, name, **_kwargs):
            nonlocal bad_attempts
            parsed_names.append(name)
            if name == "bad.pdf":
                bad_attempts += 1
                if bad_attempts == 1:
                    raise RuntimeError("simulated parse failure")
            document = Document(
                source_name=name,
                source_sha256="d" * 64,
                pages=[Page(number=1, width=612, height=792)],
            )
            rendered = render_agentic_document(document)
            return ParseResult(
                document=document,
                markdown=rendered.markdown,
                json=rendered.json,
                input_tokens=0,
                output_tokens=0,
                annotated_pdf=data,
            )

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "DocumentParser", FakeParser)
    monkeypatch.setattr(
        DocumentAgent, "analyze", lambda self, *args, **kwargs: AgenticAnalysis()
    )

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    uploader = app.file_uploader[0]
    uploader.upload("good.pdf", simple_pdf, "application/pdf")
    uploader.upload("bad.pdf", simple_pdf + b"\n", "application/pdf")
    app = uploader.run(timeout=20)
    app = _select_scanned(app, "good.pdf")
    app = _select_scanned(app, "bad.pdf")
    app = next(
        button for button in app.button if button.label == "Process 2 documents"
    ).click().run(timeout=20)

    statuses = [
        workspace["status"]
        for workspace in app.session_state["batch_workspaces"].values()
    ]
    assert not app.exception
    assert statuses == ["complete", "failed"], [
        workspace["error"]
        for workspace in app.session_state["batch_workspaces"].values()
    ]
    assert parsed_names == ["good.pdf", "bad.pdf"]

    app = next(
        button for button in app.button if button.label == "Process 2 documents"
    ).click().run(timeout=20)
    statuses = [
        workspace["status"]
        for workspace in app.session_state["batch_workspaces"].values()
    ]
    assert not app.exception
    assert statuses == ["complete", "complete"]
    assert parsed_names == ["good.pdf", "bad.pdf", "bad.pdf"]


def test_mixed_pdf_prefills_routes_and_requires_confirmation(monkeypatch) -> None:
    source = pymupdf.open()
    source.new_page().insert_text((72, 72), "Native")
    source.new_page()
    data = source.tobytes()
    source.close()
    monkeypatch.setattr(
        universal,
        "inspect_pdf_content",
        lambda _data: PdfInspection(
            pdf_type="mixed",
            page_count=2,
            pages_needing_ocr=frozenset({2}),
        ),
    )

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app = app.file_uploader[0].upload(
        "mixed.pdf", data, "application/pdf"
    ).run(timeout=20)
    selector = next(
        item
        for item in app.selectbox
        if item.label == "Processing type · mixed.pdf"
    )
    app = selector.select("Mixed PDF").run(timeout=20)

    parse = next(button for button in app.button if button.label == "Parse document")
    assert parse.disabled is True
    confirmation = next(
        item for item in app.checkbox if item.label == "Confirm page routes"
    )
    app = confirmation.check().run(timeout=20)
    parse = next(button for button in app.button if button.label == "Parse document")
    assert parse.disabled is False


def test_native_v5_result_uses_native_tabs_without_v4_agents(
    monkeypatch, simple_pdf: bytes
) -> None:
    result = NativeParseResult(
        document=NativeDocument(
            source_name="notice.pdf",
            source_sha256="a" * 64,
            source_format=SourceFormat.PDF,
            requested_processing_type=ProcessingType.NATIVE_PDF,
            base_text="Native notice",
            units=[
                SourceUnit(
                    id="page-1",
                    kind="page",
                    index=1,
                    requested_route=PageRoute.NATIVE,
                    effective_route=PageRoute.NATIVE,
                    parser="pdf-inspector",
                )
            ],
            elements=[],
        ),
        markdown="Native notice",
        json='{"schema_version":"5.0.0"}',
        annotated_pdf=simple_pdf,
    )

    class FakeUniversalParser:
        def __init__(self, *_args, **_kwargs):
            pass

        def parse(self, *_args, **_kwargs):
            return result

    monkeypatch.setattr(universal, "UniversalDocumentParser", FakeUniversalParser)
    monkeypatch.setattr(
        DocumentAgent,
        "prepare",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("v4 agent must not receive v5 results")
        ),
    )

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app = app.file_uploader[0].upload(
        "notice.pdf", simple_pdf, "application/pdf"
    ).run(timeout=20)
    selector = next(
        item
        for item in app.selectbox
        if item.label == "Processing type · notice.pdf"
    )
    app = selector.select("Native PDF").run(timeout=20)
    app = next(
        button for button in app.button if button.label == "Parse document"
    ).click().run(timeout=20)

    assert not app.exception
    assert [tab.label for tab in app.tabs][-6:] == [
        "Overview",
        "Markdown",
        "Annotated PDF",
        "Extract",
        "JSON",
        "Source Structure",
    ]
    assert any(tab.label == "Extract" for tab in app.tabs)
    assert app.session_state["result"] is result


def test_nonvisual_native_result_has_json_and_source_structure_without_pdf_tab(
    monkeypatch,
) -> None:
    source = b"<html><body><p>Native text</p></body></html>"
    result = NativeParseResult(
        document=NativeDocument(
            source_name="notice.html",
            source_sha256="b" * 64,
            source_format=SourceFormat.HTML,
            requested_processing_type=ProcessingType.OTHER_NATIVE,
            base_text="Native text",
            units=[
                SourceUnit(
                    id="document-1",
                    kind="document",
                    index=1,
                    requested_route=PageRoute.NATIVE,
                    effective_route=PageRoute.NATIVE,
                    parser="docling",
                )
            ],
            elements=[],
        ),
        markdown="Native text",
        json='{"schema_version":"5.0.0"}',
    )

    class FakeUniversalParser:
        def __init__(self, *_args, **_kwargs):
            pass

        def parse(self, *_args, **_kwargs):
            return result

    monkeypatch.setattr(universal, "UniversalDocumentParser", FakeUniversalParser)

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app = app.file_uploader[0].upload(
        "notice.html", source, "text/html"
    ).run(timeout=20)
    selector = next(
        item
        for item in app.selectbox
        if item.label == "Processing type · notice.html"
    )
    app = selector.select("Other Native").run(timeout=20)
    app = next(
        button for button in app.button if button.label == "Parse document"
    ).click().run(timeout=20)

    assert not app.exception
    assert [tab.label for tab in app.tabs] == [
        "Overview",
        "Markdown",
        "Extract",
        "JSON",
        "Source Structure",
    ]
    assert not any(button.label == "Download annotated PDF" for button in app.button)


def test_completed_batch_restores_after_app_restart_without_reparsing(
    monkeypatch, simple_pdf: bytes
) -> None:
    parse_calls = 0

    class FakeParser:
        def __init__(self, config=None):
            assert config.ocr_engine.value == "glm-ocr"

        def parse(self, data, name, **_kwargs):
            nonlocal parse_calls
            parse_calls += 1
            document = Document(
                source_name=name,
                source_sha256="e" * 64,
                pages=[Page(number=1, width=612, height=792)],
            )
            rendered = render_agentic_document(document)
            return ParseResult(
                document=document,
                markdown=rendered.markdown,
                json=rendered.json,
                input_tokens=0,
                output_tokens=0,
                annotated_pdf=data,
            )

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "DocumentParser", FakeParser)
    monkeypatch.setattr(
        DocumentAgent, "analyze", lambda self, *args, **kwargs: AgenticAnalysis()
    )

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app = app.file_uploader[0].upload(
        "notice.pdf", simple_pdf, "application/pdf"
    ).run(timeout=20)
    app = _select_scanned(app, "notice.pdf")
    app = next(
        button for button in app.button if button.label == "Parse document"
    ).click().run(timeout=20)
    assert not app.exception
    assert parse_calls == 1

    restarted = AppTest.from_file("streamlit_app.py").run(timeout=20)

    assert not restarted.exception
    assert parse_calls == 1
    assert restarted.session_state["result"].document.source_name == "notice.pdf"
    assert restarted.session_state["batch_workspaces"][
        restarted.session_state["active_document_id"]
    ]["status"] == "complete"
    assert any(
        button.label == "Download Markdown" for button in restarted.download_button
    )
    assert any(
        "Restored 1 document(s)" in item.value for item in restarted.get("caption")
    )


def test_interrupted_analysis_resumes_from_parse_checkpoint(
    monkeypatch, simple_pdf: bytes, tmp_path
) -> None:
    database = tmp_path / "studio.sqlite3"
    document = build_batch_documents(
        [("notice.pdf", simple_pdf, "application/pdf")]
    )[0]
    parsed_document = Document(
        source_name="notice.pdf",
        source_sha256="f" * 64,
        pages=[Page(number=1, width=612, height=792)],
    )
    rendered = render_agentic_document(parsed_document)
    result = ParseResult(
        document=parsed_document,
        markdown=rendered.markdown,
        json=rendered.json,
        input_tokens=0,
        output_tokens=0,
        annotated_pdf=simple_pdf,
    )
    selection_key = (
        f"{document.id}:all:False:True:True:glm-ocr:scanned-pdf::4.6.0"
    )
    store = WorkspaceStore(database)
    store.sync_documents(
        [document],
        settings={
            "ade_mode": "Fast",
            "refine_markdown": False,
            "classify_document": True,
            "generate_toc": False,
            "visual_recovery": True,
            "ocr_engine_label": "GLM-OCR",
        },
        result_version="4.6.0",
    )
    store.save_document(
        document.id,
        status="processing",
        selection_key=selection_key,
        parsed_source=simple_pdf,
        result=result,
        progress={
            "stage": "classify",
            "current": 0,
            "total": 1,
            "message": "Running Luna document analysis",
        },
    )
    analysis_calls = 0

    class FakeParser:
        def __init__(self, config=None):
            assert config.ocr_engine.value == "glm-ocr"

        def parse(self, *_args, **_kwargs):
            raise AssertionError("OCR checkpoint should be reused")

    def analyze(self, *args, **kwargs):
        nonlocal analysis_calls
        analysis_calls += 1
        return AgenticAnalysis()

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "DocumentParser", FakeParser)
    monkeypatch.setattr(DocumentAgent, "analyze", analyze)

    app = AppTest.from_file("streamlit_app.py").run(timeout=20)
    app = next(
        button for button in app.button if button.label == "Resume batch"
    ).click().run(timeout=20)

    assert not app.exception
    assert analysis_calls == 1
    assert app.session_state["batch_workspaces"][document.id]["status"] == "complete"
