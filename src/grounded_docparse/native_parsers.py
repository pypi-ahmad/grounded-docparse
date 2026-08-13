from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from io import BytesIO
from types import ModuleType, SimpleNamespace

import pymupdf

from .config import ParserConfig
from .models import Element, RunUsage
from .native import (
    CellSourceAnchor,
    CsvSourceAnchor,
    NativeDocument,
    NativeElement,
    NativeParseResult,
    PageRoute,
    PdfSourceAnchor,
    ProcessingType,
    SourceFormat,
    SourceSpan,
    SourceUnit,
    StructuralSourceAnchor,
    render_native_document,
)
from .pipeline import DocumentParser
from .render import build_elements, render_agentic_document, render_annotated_pdf
from .universal import (
    MixedNativePageUnusable,
    NativePdfRequiresMixed,
    PdfInspection,
)


def _pdf_inspector() -> ModuleType:
    try:
        import pdf_inspector
    except ImportError as exc:
        raise RuntimeError(
            "native PDF parsing requires grounded-docparse[native]"
        ) from exc
    return pdf_inspector


def _subset_pdf(data: bytes, pages: list[int]) -> bytes:
    with pymupdf.open(stream=data, filetype="pdf") as source:
        output = pymupdf.open()
        try:
            for page in pages:
                output.insert_pdf(source, from_page=page - 1, to_page=page - 1)
            return output.tobytes(garbage=3, deflate=True)
        finally:
            output.close()


def _native_type(item, roles: dict[tuple[int, int], str]) -> str:
    role = roles.get((item.page, item.mcid)) if item.mcid is not None else None
    if role:
        return role.casefold()
    return str(item.item_type or "text").casefold()


class PdfInspectorParser:
    def __init__(
        self,
        config: ParserConfig,
        legacy_parser: DocumentParser,
        *,
        pdf_module: ModuleType | None = None,
    ) -> None:
        self.config = config
        self.legacy_parser = legacy_parser
        self.pdf_module = pdf_module or _pdf_inspector()

    def parse(
        self,
        data: bytes,
        filename: str,
        *,
        processing_type: ProcessingType,
        page_routes: dict[int, PageRoute] | None,
        inspection: PdfInspection,
        progress_callback=None,
        refine_markdown: bool = True,
        visual_recovery: bool = True,
    ) -> NativeParseResult:
        routes = (
            {page: PageRoute.NATIVE for page in range(1, inspection.page_count + 1)}
            if processing_type is ProcessingType.NATIVE_PDF
            else dict(page_routes or {})
        )
        native_pages = [page for page, route in routes.items() if route is PageRoute.NATIVE]
        ocr_pages = [page for page, route in routes.items() if route is PageRoute.OCR]
        unusable_native = set(native_pages) & set(inspection.pages_needing_ocr)
        if unusable_native:
            if processing_type is ProcessingType.NATIVE_PDF:
                raise NativePdfRequiresMixed(unusable_native)
            raise MixedNativePageUnusable(unusable_native)

        page_markdown: dict[int, str] = {}
        native_items = []
        roles: dict[tuple[int, int], str] = {}
        if native_pages:
            extracted = self.pdf_module.extract_pages_markdown_bytes(
                data, pages=[page - 1 for page in native_pages]
            )
            unusable = {
                page.page + 1 for page in extracted.pages if page.needs_ocr
            }
            if unusable:
                if processing_type is ProcessingType.NATIVE_PDF:
                    raise NativePdfRequiresMixed(unusable)
                raise MixedNativePageUnusable(unusable)
            page_markdown.update(
                (page.page + 1, page.markdown) for page in extracted.pages
            )
            native_items = self.pdf_module.extract_text_with_positions_bytes(
                data, pages=native_pages
            )
            roles = {
                (item.page, item.mcid): item.role
                for item in self.pdf_module.extract_structure_elements_bytes(
                    data, pages=native_pages
                )
            }

        ocr_result = None
        ocr_page_map: dict[int, int] = {}
        if ocr_pages:
            subset = _subset_pdf(data, ocr_pages)
            ocr_result = self.legacy_parser.parse(
                subset,
                filename,
                progress_callback=progress_callback,
                refine_markdown=refine_markdown,
                visual_recovery=visual_recovery,
            )
            ocr_page_map = {
                subset_page: source_page
                for subset_page, source_page in enumerate(ocr_pages, start=1)
            }
            for page in ocr_result.document.pages:
                source_page = ocr_page_map[page.number]
                copied = page.model_copy(deep=True, update={"number": source_page})
                rendered = render_agentic_document(
                    ocr_result.document.model_copy(update={"pages": [copied]})
                )
                page_markdown[source_page] = rendered.markdown

        with pymupdf.open(stream=data, filetype="pdf") as source:
            page_sizes = {
                index + 1: (page.rect.width, page.rect.height)
                for index, page in enumerate(source)
            }

        base_parts: list[str] = []
        elements: list[NativeElement] = []
        annotation_elements: list[Element] = []
        reading_order: defaultdict[int, int] = defaultdict(int)
        base_length = 0
        candidates: defaultdict[
            int, list[tuple[int, str, str, tuple[float, float, float, float] | None]]
        ] = defaultdict(list)

        def add_element(
            *, page: int, text: str, kind: str, bbox: tuple[float, float, float, float] | None
        ) -> None:
            nonlocal base_length
            if not text:
                return
            if base_parts:
                base_parts.append("\n")
                base_length += 1
            start = base_length
            base_parts.append(text)
            end = start + len(text)
            base_length = end
            reading_order[page] += 1
            element_id = f"p{page}-e{reading_order[page]}"
            anchor = PdfSourceAnchor(
                unit_id=f"page-{page}", page=page, bbox=bbox
            )
            elements.append(
                NativeElement(
                    id=element_id,
                    type=kind,
                    text=text,
                    reading_order=reading_order[page] - 1,
                    source=SourceSpan(
                        start=start, end=end, element_id=element_id, anchor=anchor
                    ),
                )
            )
            annotation_elements.append(
                Element(
                    id=element_id,
                    type=kind,
                    page=page,
                    bbox=bbox,
                    text=text,
                    reading_order=reading_order[page],
                )
            )

        for item in native_items:
            width, height = page_sizes[item.page]
            bbox = (
                max(0.0, item.x / width),
                max(0.0, item.y / height),
                min(1.0, (item.x + item.width) / width),
                min(1.0, (item.y + item.height) / height),
            )
            candidates[item.page].append(
                (len(candidates[item.page]), item.text, _native_type(item, roles), bbox)
            )

        if ocr_result is not None:
            ocr_elements = ocr_result.elements or build_elements(ocr_result.document)
            for item in sorted(ocr_elements, key=lambda value: (value.page, value.reading_order)):
                source_page = ocr_page_map[item.page]
                candidates[source_page].append(
                    (item.reading_order, item.text, str(item.type), item.bbox)
                )

        for page in range(1, inspection.page_count + 1):
            for _order, text, kind, bbox in sorted(candidates[page]):
                add_element(page=page, text=text, kind=kind, bbox=bbox)

        units = []
        for page in range(1, inspection.page_count + 1):
            warnings = []
            if page in inspection.pages_with_tables:
                warnings.append("table layout detected")
            if page in inspection.pages_with_columns:
                warnings.append("multi-column layout detected")
            units.append(
                SourceUnit(
                    id=f"page-{page}",
                    kind="page",
                    index=page,
                    requested_route=routes[page],
                    effective_route=routes[page],
                    parser=(
                        "pdf-inspector"
                        if routes[page] is PageRoute.NATIVE
                        else self.config.ocr_engine.value
                    ),
                    warnings=warnings,
                )
            )
        document = NativeDocument(
            source_name=filename,
            source_sha256=hashlib.sha256(data).hexdigest(),
            source_format=SourceFormat.PDF,
            requested_processing_type=processing_type,
            base_text="".join(base_parts),
            units=units,
            elements=elements,
        )
        markdown = "\n\n".join(
            f"<!-- Page {page} -->\n\n{page_markdown.get(page, '').strip()}"
            for page in range(1, inspection.page_count + 1)
        ).rstrip()
        rendered = render_native_document(document, markdown=markdown)
        annotated_pdf = render_annotated_pdf(
            data,
            filename,
            annotation_elements,
            page_count=inspection.page_count,
        )
        usage = (
            ocr_result.usage.model_copy(deep=True)
            if ocr_result is not None and ocr_result.usage is not None
            else RunUsage()
        )
        return NativeParseResult(
            document=document,
            markdown=rendered.markdown,
            json=rendered.json,
            annotated_pdf=annotated_pdf,
            input_tokens=usage.input_tokens,
            output_tokens=usage.output_tokens,
            usage=usage,
            trace=list(ocr_result.trace or []) if ocr_result is not None else [],
        )


class DoclingNativeParser:
    def __init__(self, config: ParserConfig, *, converter=None) -> None:
        self.config = config
        self.converter = converter

    def parse(
        self,
        data: bytes,
        filename: str,
        *,
        source_format: SourceFormat,
        processing_type: ProcessingType,
    ) -> NativeParseResult:
        from .docling_native import (
            DOCLING_FORMAT_NAMES,
            build_source_manifest,
            claim_record,
            make_docling_converter,
        )

        if source_format not in DOCLING_FORMAT_NAMES:
            raise ValueError(f"Docling native parsing does not support {source_format.value}")
        manifest = build_source_manifest(data, source_format)
        converter = self.converter or make_docling_converter()
        self.converter = converter
        try:
            from docling.datamodel.document import DocumentStream
            from docling_core.types.doc import TableItem, TextItem
        except ImportError as exc:
            raise RuntimeError(
                "native document parsing requires grounded-docparse[native]"
            ) from exc
        result = converter.convert(
            DocumentStream(name=filename, stream=BytesIO(data)),
            raises_on_error=True,
            max_file_size=self.config.max_upload_bytes,
        )
        docling_document = result.document
        base_parts: list[str] = []
        elements: list[NativeElement] = []
        base_length = 0

        def append_text(value: str) -> tuple[int, int]:
            nonlocal base_length
            if base_parts:
                base_parts.append("\n")
                base_length += 1
            start = base_length
            base_parts.append(value)
            base_length += len(value)
            return start, base_length

        def unit_for(item) -> str | None:
            provenance = getattr(item, "prov", None) or []
            if provenance and 1 <= provenance[0].page_no <= len(manifest.units):
                return manifest.units[provenance[0].page_no - 1].id
            return None

        for item, _level in docling_document.iterate_items(with_groups=False):
            if isinstance(item, TextItem):
                parent_ref = getattr(getattr(item, "parent", None), "cref", "")
                if str(parent_ref).startswith("#/tables/"):
                    continue
                value = item.text.strip()
                if not value:
                    continue
                if any(
                    value in {asset.alt_text, asset.caption}
                    for asset in manifest.assets
                ):
                    continue
                try:
                    record = claim_record(
                        manifest.records, value, unit_id=unit_for(item)
                    )
                except ValueError:
                    matching_unit = next(
                        (
                            unit
                            for unit in manifest.units
                            if unit.label
                            and " ".join(unit.label.split()).casefold()
                            == " ".join(value.split()).casefold()
                        ),
                        None,
                    )
                    if source_format is not SourceFormat.ODP or matching_unit is None:
                        raise
                    record = SimpleNamespace(
                        anchor=StructuralSourceAnchor(
                            unit_id=matching_unit.id,
                            path=(
                                f"/office:presentation/draw:page[{matching_unit.index}]"
                                "/@draw:name"
                            ),
                        )
                    )
                start, end = append_text(value)
                element_id = f"element-{len(elements) + 1}"
                elements.append(
                    NativeElement(
                        id=element_id,
                        type=str(item.label.value),
                        text=value,
                        reading_order=len(elements),
                        source=SourceSpan(
                            start=start,
                            end=end,
                            element_id=element_id,
                            anchor=record.anchor,
                        ),
                    )
                )
            elif isinstance(item, TableItem):
                cells = sorted(
                    item.data.table_cells,
                    key=lambda cell: (
                        cell.start_row_offset_idx,
                        cell.start_col_offset_idx,
                    ),
                )
                record = claim_record(
                    manifest.records,
                    " ".join(cell.text for cell in cells),
                    unit_id=unit_for(item),
                    record_type="table",
                )
                if record.cells is not None:
                    cells = [
                        SimpleNamespace(
                            text=value,
                            start_row_offset_idx=row,
                            end_row_offset_idx=row + 1,
                            start_col_offset_idx=column,
                            end_col_offset_idx=column + 1,
                        )
                        for row, column, value in record.cells
                    ]
                    row_count = max((cell.start_row_offset_idx for cell in cells), default=-1) + 1
                    column_count = max((cell.start_col_offset_idx for cell in cells), default=-1) + 1
                else:
                    row_count = item.data.num_rows
                    column_count = item.data.num_cols
                rows = [["" for _ in range(column_count)] for _ in range(row_count)]
                for cell in cells:
                    rows[cell.start_row_offset_idx][cell.start_col_offset_idx] = cell.text
                table_text = "\n".join("\t".join(row) for row in rows)
                start, end = append_text(table_text)
                table_id = f"element-{len(elements) + 1}"
                table_element = NativeElement(
                    id=table_id,
                    type="table",
                    text=table_text,
                    reading_order=len(elements),
                    source=SourceSpan(
                        start=start,
                        end=end,
                        element_id=table_id,
                        anchor=record.anchor,
                    ),
                )
                elements.append(table_element)
                cursor = start
                for cell in cells:
                    value = cell.text
                    if not value:
                        continue
                    cell_start = "".join(base_parts).find(value, cursor, end)
                    if cell_start < 0:
                        raise ValueError("table cell cannot be mapped into immutable base_text")
                    cursor = cell_start + len(value)
                    cell_id = f"element-{len(elements) + 1}"
                    anchor = record.anchor
                    if isinstance(anchor, CellSourceAnchor):
                        match = re.match(r"([A-Z]+)(\d+)", anchor.cell_range)
                        if match is None:
                            raise ValueError("spreadsheet table has an invalid source range")
                        from openpyxl.utils.cell import (
                            column_index_from_string,
                            get_column_letter,
                        )

                        base_column = column_index_from_string(match.group(1))
                        base_row = int(match.group(2))
                        cell_anchor = CellSourceAnchor(
                            unit_id=anchor.unit_id,
                            sheet=anchor.sheet,
                            cell_range=(
                                f"{get_column_letter(base_column + cell.start_col_offset_idx)}"
                                f"{base_row + cell.start_row_offset_idx}"
                            ),
                        )
                    elif isinstance(anchor, CsvSourceAnchor):
                        cell_anchor = CsvSourceAnchor(
                            unit_id=anchor.unit_id,
                            row_start=anchor.row_start + cell.start_row_offset_idx,
                            row_end=anchor.row_start + cell.end_row_offset_idx - 1,
                            column_start=anchor.column_start + cell.start_col_offset_idx,
                            column_end=anchor.column_start + cell.end_col_offset_idx - 1,
                        )
                    else:
                        cell_anchor = StructuralSourceAnchor(
                            unit_id=anchor.unit_id,
                            path=(
                                f"{anchor.path}/row[{cell.start_row_offset_idx + 1}]"
                                f"/cell[{cell.start_col_offset_idx + 1}]"
                            ),
                            bbox=getattr(anchor, "bbox", None),
                        )
                    elements.append(
                        NativeElement(
                            id=cell_id,
                            type="table_cell",
                            text=value,
                            reading_order=len(elements),
                            source=SourceSpan(
                                start=cell_start,
                                end=cursor,
                                element_id=cell_id,
                                anchor=cell_anchor,
                            ),
                        )
                    )
                    table_element.children.append(cell_id)

        unclaimed = [record for record in manifest.records if record.text and not record.claimed]
        if unclaimed:
            raise ValueError(
                "Docling omitted source block(s): "
                + ", ".join(record.anchor.model_dump_json() for record in unclaimed[:3])
            )
        document = NativeDocument(
            source_name=filename,
            source_sha256=hashlib.sha256(data).hexdigest(),
            source_format=source_format,
            requested_processing_type=processing_type,
            base_text="".join(base_parts),
            units=manifest.units,
            elements=elements,
            assets=manifest.assets,
        )
        child_ids = {child for element in elements for child in element.children}
        markdown_parts = []
        for element in elements:
            if element.id in child_ids:
                continue
            if element.type == "table":
                rows = [row.split("\t") for row in element.text.splitlines()]
                width = max((len(row) for row in rows), default=0)
                if width:
                    normalized = [row + [""] * (width - len(row)) for row in rows]
                    escape = lambda value: value.replace("|", "\\|").replace("\n", " ")
                    markdown_parts.append(
                        "\n".join(
                            [
                                "| " + " | ".join(escape(value) for value in normalized[0]) + " |",
                                "| " + " | ".join("---" for _ in range(width)) + " |",
                            ]
                            + [
                                "| " + " | ".join(escape(value) for value in row) + " |"
                                for row in normalized[1:]
                            ]
                        )
                    )
            elif element.type == "title":
                markdown_parts.append(f"# {element.text}")
            elif element.type == "section_header":
                markdown_parts.append(f"## {element.text}")
            elif element.type == "list_item":
                markdown_parts.append(f"- {element.text}")
            else:
                markdown_parts.append(element.text)
        markdown = "\n\n".join(markdown_parts)
        rendered = render_native_document(document, markdown=markdown)
        return NativeParseResult(
            document=document,
            markdown=rendered.markdown,
            json=rendered.json,
        )
